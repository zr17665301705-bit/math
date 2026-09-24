"""Create the detailed Word chapter for the coupled PEMFC cold-start model."""
from pathlib import Path
import math

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUTDIR = Path("model_outputs/Q1_cold_start_coupled_v2")
OUT = OUTDIR / "问题1模型建立_详细耦合优化版.docx"
FLOW = OUTDIR / "模型耦合求解流程.png"


def draw_flowchart():
    W,H=2016,972; img=Image.new("RGB",(W,H),"white"); dr=ImageDraw.Draw(img)
    navy,blue,pale,edge="#183B63","#2E6F9E","#EAF2F8","#AEBBC7"
    font_path=r"C:\Windows\Fonts\msyh.ttc"; bold_path=r"C:\Windows\Fonts\msyhbd.ttc"
    f_title=ImageFont.truetype(bold_path,29); f_body=ImageFont.truetype(font_path,23); f_small=ImageFont.truetype(font_path,20); f_note=ImageFont.truetype(bold_path,24)

    def center_text(xy,text,font,fill,spacing=6):
        box=dr.multiline_textbbox((0,0),text,font=font,spacing=spacing,align="center")
        tw,th=box[2]-box[0],box[3]-box[1]; x=(xy[0]+xy[2]-tw)/2; y=(xy[1]+xy[3]-th)/2
        dr.multiline_text((x,y),text,font=font,fill=fill,spacing=spacing,align="center")

    def node(x,y,w,h,title,body):
        dr.rounded_rectangle((x,y,x+w,y+h),radius=18,fill=pale,outline=edge,width=3)
        dr.text((x+w/2,y+35),title,font=f_title,fill=navy,anchor="mm")
        center_text((x+8,y+65,x+w-8,y+h-8),body,f_body,"#263746")

    def arrow(points,label=""):
        dr.line(points,fill=blue,width=4,joint="curve")
        x2,y2=points[-1]; x1,y1=points[-2]; ang=math.atan2(y2-y1,x2-x1); size=16
        p1=(x2-size*math.cos(ang-.55),y2-size*math.sin(ang-.55)); p2=(x2-size*math.cos(ang+.55),y2-size*math.sin(ang+.55))
        dr.polygon([(x2,y2),p1,p2],fill=blue)
        if label:
            mx=(points[0][0]+points[-1][0])/2; my=(points[0][1]+points[-1][1])/2-16
            dr.text((mx,my),label,font=f_small,fill=blue,anchor="mm")

    boxes=[(40,80,295,220,"外部输入","j(t),  T_amb\n初始状态"),(390,80,315,220,"电荷与膜状态","q,  lambda(q)\n失水记忆 z"),
           (765,80,325,220,"阴极水相","m_v, m_l, m_i\n凝结 冻结 融化"),(1150,80,325,220,"冰堵塞修正","s_ice → a/a0\neps_eff → D_O2,eff"),
           (1535,80,400,220,"电化学电压","E_rev − eta_act − eta_ohm\n− eta_con − eta_dry"),(1320,570,360,210,"热源","j(Eth−V)\n+ Lf dm_i/dt"),
           (810,570,410,210,"七层温度场","隐式有限体积\nT1 … T7,  T_mean")]
    for b in boxes: node(*b)
    arrow([(335,190),(390,190)],"j"); arrow([(705,190),(765,190)],"j, q"); arrow([(1090,190),(1150,190)],"m_i")
    arrow([(1475,190),(1535,190)],"j0, D_O2"); arrow([(1735,300),(1735,465),(1500,570)],"V, j")
    arrow([(1320,675),(1220,675)],"Q"); arrow([(810,675),(630,675),(630,300),(545,300)],"T_mean → kappa, j0")
    arrow([(1015,570),(1015,300)],"T_mean → p_sat, 相变"); arrow([(185,300),(185,850),(810,850)],"T_amb")
    dr.text((W/2,925),"每个时间步内重复  水相 → 冰堵塞 → 电压 → 热源 → 温度  直至 T_mean 与 V 收敛",font=f_note,fill=navy,anchor="mm")
    img.save(FLOW,dpi=(180,180))


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tcPr.append(shd)


def set_cell_margins(cell, v=70, h=90):
    tcPr = cell._tc.get_or_add_tcPr(); mar = OxmlElement("w:tcMar")
    for side, width in (("top",v),("bottom",v),("left",h),("right",h)):
        e=OxmlElement("w:"+side); e.set(qn("w:w"),str(width)); e.set(qn("w:type"),"dxa"); mar.append(e)
    tcPr.append(mar)


def set_borders(cell, color="D9D9D9"):
    tcPr=cell._tc.get_or_add_tcPr(); borders=OxmlElement("w:tcBorders")
    for side in ("top","left","bottom","right","insideH","insideV"):
        e=OxmlElement("w:"+side); e.set(qn("w:val"),"single"); e.set(qn("w:sz"),"4"); e.set(qn("w:color"),color); borders.append(e)
    tcPr.append(borders)


def repeat_header(row):
    trPr=row._tr.get_or_add_trPr(); e=OxmlElement("w:tblHeader"); e.set(qn("w:val"),"true"); trPr.append(e)


def add_table(doc, headers, rows, widths, aligns=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    repeat_header(t.rows[0])
    for i,(h,w) in enumerate(zip(headers,widths)):
        t.columns[i].width=Inches(w); t.rows[0].cells[i].text=h
    for row in rows:
        cells=t.add_row().cells
        for i,val in enumerate(row): cells[i].text=str(val)
    for ri,row in enumerate(t.rows):
        for ci,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell); set_borders(cell)
            set_cell_shading(cell,"183B63" if ri==0 else ("F2F6FA" if ri%2==0 else "FFFFFF"))
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.1
                if aligns: p.alignment=aligns[ci]
                for r in p.runs:
                    r.font.size=Pt(9.2); r.font.name="宋体"; r._element.rPr.rFonts.set(qn("w:eastAsia"),"宋体")
                    if ri==0: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t


def add_eq(doc, text, number=None):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(4)
    ompara=OxmlElement("m:oMathPara"); om=OxmlElement("m:oMath"); mr=OxmlElement("m:r")
    mrpr=OxmlElement("m:rPr"); sty=OxmlElement("m:sty"); sty.set(qn("m:val"),"p"); mrpr.append(sty); mr.append(mrpr)
    mt=OxmlElement("m:t"); mt.text=text + (f"    ({number})" if number else ""); mr.append(mt); om.append(mr); ompara.append(om); p._p.append(ompara)
    return p


def add_para(doc, text, first=True, bold_lead=None):
    p=doc.add_paragraph()
    if first: p.paragraph_format.first_line_indent=Pt(21)
    if bold_lead and text.startswith(bold_lead):
        r=p.add_run(bold_lead); r.bold=True; p.add_run(text[len(bold_lead):])
    else: p.add_run(text)
    return p


def add_bullet(doc, text):
    p=doc.add_paragraph(style="List Bullet"); p.paragraph_format.left_indent=Pt(21); p.paragraph_format.first_line_indent=Pt(-10); p.add_run(text)


def build():
    OUTDIR.mkdir(parents=True, exist_ok=True); draw_flowchart()
    d=Document(); sec=d.sections[0]
    sec.page_width,sec.page_height=Inches(8.5),Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.72); sec.left_margin=sec.right_margin=Inches(.82)
    styles=d.styles
    for name in ("Normal","Title","Heading 1","Heading 2","Heading 3"):
        s=styles[name]; font="宋体" if name=="Normal" else "黑体"
        s.font.name=font; s._element.rPr.rFonts.set(qn("w:eastAsia"),font); s.font.color.rgb=RGBColor(0,0,0)
    styles["Normal"].font.size=Pt(10.5); styles["Normal"].paragraph_format.line_spacing=1.22; styles["Normal"].paragraph_format.space_after=Pt(5)
    styles["Title"].font.size=Pt(17); styles["Title"].font.bold=True; styles["Title"].paragraph_format.space_after=Pt(10)
    styles["Heading 1"].font.size=Pt(13); styles["Heading 1"].font.bold=True; styles["Heading 1"].paragraph_format.space_before=Pt(12); styles["Heading 1"].paragraph_format.space_after=Pt(5); styles["Heading 1"].paragraph_format.keep_with_next=True
    styles["Heading 2"].font.size=Pt(11.3); styles["Heading 2"].font.bold=True; styles["Heading 2"].paragraph_format.space_before=Pt(8); styles["Heading 2"].paragraph_format.space_after=Pt(3); styles["Heading 2"].paragraph_format.keep_with_next=True
    styles["Heading 3"].font.size=Pt(10.5); styles["Heading 3"].font.bold=True; styles["Heading 3"].paragraph_format.space_before=Pt(6); styles["Heading 3"].paragraph_format.space_after=Pt(2); styles["Heading 3"].paragraph_format.keep_with_next=True
    # Remove Word's theme title rule.
    ppr=styles["Title"]._element.get_or_add_pPr()
    for e in list(ppr):
        if e.tag==qn("w:pBdr"): ppr.remove(e)

    title=d.add_paragraph("问题一  PEMFC 低温冷启动详细耦合模型",style="Title"); title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    sub=d.add_paragraph("模型建立  数值求解  参数辨识  留出验证")
    sub.alignment=WD_ALIGN_PARAGRAPH.CENTER; sub.paragraph_format.space_after=Pt(12)
    for r in sub.runs: r.font.size=Pt(10); r.font.color.rgb=RGBColor(70,70,70)
    add_para(d,"本文根据附件 1 的电池结构与热物性、附件 2 的 −20°C 和 −25°C 冷启动序列，建立沿膜电极厚度方向的一维瞬态模型。模型同时求解七层温度、阴极水蒸气与液水及冰库存、膜水化、阳极暂态失水和端电压。新版的核心改进是按照 Jiao 和 Li 的多相模型，将冰覆盖对交换电流源项的影响由过强的幂函数改为未堵塞比例的一次函数，并在每个时间步内迭代水相变、电压、热源和温度，使双向耦合在程序中真正闭合。")
    add_para(d,"两种温度均取前 70% 数据辨识参数，后 30% 完全留出。新版在留出段的电压 RMSE 分别为 0.01990 V 和 0.01202 V，比上一版降低 26.2% 和 42.6%。水冰库存缺少直接观测，因此端电压和平均温度的吻合只验证可观测输出，不能等同于冰分布已经得到实验验证。")

    d.add_heading("1 模型目标与计算域",level=1)
    add_para(d,"模型需要同时解释三个现象：升载后电压的快速下降、产水使膜逐渐水化带来的电压恢复，以及冰累积对反应面积和氧传质的抑制。温度既受反应热和凝固潜热驱动，又反过来改变膜电导、交换电流、氧扩散和相变状态，因此必须把质量守恒、电化学和能量守恒联立求解。")
    d.add_heading("1.1 一维七层结构",level=2)
    add_para(d,"沿厚度方向依次布置阳极双极板、阳极气体扩散层、阳极催化层、质子交换膜、阴极催化层、阴极气体扩散层和阴极双极板。每层设一个有限体积控制单元。反应产水、反应热和凝固潜热集中施加于阴极催化层；固体层间满足温度和热流连续。")
    add_table(d,["层号","控制体","厚度","热参数来源"],[
        (1,"阳极双极板","2.000 mm","附件 1"),(2,"阳极 GDL","0.150 mm","附件 1"),(3,"阳极 CL","3.4 μm","附件 1"),
        (4,"PEM","12.0 μm","附件 1"),(5,"阴极 CL","11.3 μm","附件 1"),(6,"阴极 GDL","0.150 mm","附件 1"),(7,"阴极双极板","2.000 mm","附件 1")
    ],[.55,1.65,1.05,2.5],[WD_ALIGN_PARAGRAPH.CENTER]*4)
    d.add_heading("1.2 主要假设",level=2)
    for x in [
        "短时冷启动期间入口压力和氧摩尔分数保持不变，附件 2 的电流密度作为外部输入。",
        "阴极催化层用单位几何面积上的三个集中库存描述水蒸气、液水和冰；没有足够观测把水相进一步分辨到各层和面内位置。",
        "初始温度取每个工况的首个实测平均温度；初始液水、冰和暂态失水记忆均为零，初始膜含水量 λ₀=3。",
        "水滞留率和冻结速率由先验给定；交换电流、传质倍率、膜水化和失水电压系数作为本实验条件下的有效参数。"
    ]: add_bullet(d,x)

    d.add_heading("2 状态变量与方程耦合关系",level=1)
    add_para(d,"时刻 t 的模型状态写为下式。温度向量决定相变、膜电导和电化学动力学；水冰状态决定孔隙与反应面积；电压又决定反应放热，形成闭环。")
    add_eq(d,"X(t) = [T₁,…,T₇, mᵥ, mₗ, mᵢ, q, λ, z]ᵀ",1)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(2)
    run=p.add_run(); run.add_picture(str(FLOW),width=Inches(6.75))
    cap=d.add_paragraph("图 1  多物理方程的双向耦合与步内迭代顺序"); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER; cap.paragraph_format.space_after=Pt(7)
    for r in cap.runs: r.font.size=Pt(9); r.font.bold=True
    add_table(d,["被更新量","直接输入","向后影响"],[
        ("累计电荷 q","电流密度 j","膜含水量 λ、失水激励衰减"),
        ("水相 mᵥ mₗ mᵢ","j、温度 T̄","冰饱和度、潜热"),
        ("膜状态 λ z","q、电流增量","欧姆损失、暂态失水损失"),
        ("冰堵塞 sᵢ","冰质量 mᵢ","交换电流、氧扩散、浓差损失"),
        ("端电压 V","T̄、j、λ、z、sᵢ","反应热 j(Eth−V)"),
        ("温度 T₁…T₇","反应热、潜热、环境散热","相变、动力学、传质与电压")
    ],[1.25,1.75,3.25],[WD_ALIGN_PARAGRAPH.CENTER,WD_ALIGN_PARAGRAPH.LEFT,WD_ALIGN_PARAGRAPH.LEFT])

    d.add_heading("3 七层瞬态热模型",level=1)
    d.add_heading("3.1 连续方程与边界条件",level=2)
    add_para(d,"第 i 层满足一维非稳态导热方程。除阴极催化层外，体热源为零；阴极催化层的面热源通过控制体厚度换算到源项。")
    add_eq(d,"ρᵢ cₚ,ᵢ ∂Tᵢ/∂t = ∂/∂x (kᵢ ∂Tᵢ/∂x) + Sᵀ,ᵢ",2)
    add_eq(d,"Tᵢ = Tᵢ₊₁,   kᵢ ∂Tᵢ/∂x = kᵢ₊₁ ∂Tᵢ₊₁/∂x",3)
    add_eq(d,"−k ∂T/∂n = h(T−Tₐₘᵦ),   h = 40 W·m⁻²·K⁻¹",4)
    d.add_heading("3.2 有限体积离散",level=2)
    add_para(d,"第 i 个控制体的单位面积热容为 Cᵢ=sCρᵢcₚ,ᵢΔxᵢ。相邻中心之间的热导采用两侧半层热阻串联，避免材料界面直接平均导热率。时间项采用隐式欧拉格式。")
    add_eq(d,"Gᵢ₊½ = [Δxᵢ/(2kᵢ)+Δxᵢ₊₁/(2kᵢ₊₁)]⁻¹",5)
    add_eq(d,"Cᵢ(Tᵢⁿ⁺¹−Tᵢⁿ)/Δt = Gᵢ₋½(Tᵢ₋₁ⁿ⁺¹−Tᵢⁿ⁺¹)+Gᵢ₊½(Tᵢ₊₁ⁿ⁺¹−Tᵢⁿ⁺¹)+Qᵢ",6)
    add_para(d,"模型输出的平均温度按控制体厚度加权。上一版早期程序曾对厚度差异很大的七个节点直接算术平均，会夸大微米级催化层和膜的权重；该问题已经修正。")
    add_eq(d,"T̄ = Σᵢ ΔxᵢTᵢ / Σᵢ Δxᵢ",7)

    d.add_heading("4 阴极产水与水冰相变",level=1)
    d.add_heading("4.1 法拉第产水与蒸气容量",level=2)
    add_para(d,"单位几何面积上的法拉第产水速率为 jMw/(2F)。设 r=0.12 的产水留在阴极催化层，其余随气流排出。保留水先进入蒸气库存；当蒸气超过孔隙在当前温度下的饱和容量时，超出部分转入液态库存。")
    add_eq(d,"ṁprod = r j Mw/(2F)",8)
    add_eq(d,"mᵥ,sat = ε₀ LCCL Mw pₛₐₜ(T)/(RT)",9)
    add_para(d,"代码在 0°C 以上使用水面饱和蒸气压，在 0°C 以下使用冰面饱和蒸气压。温度用摄氏值 Tc 代入：")
    add_eq(d,"pₛₐₜ = 611.21 exp[(18.678−Tc/234.5)Tc/(257.14+Tc)],   Tc ≥ 0",10)
    add_eq(d,"pₛₐₜ = 611.15 exp[(23.036−Tc/333.7)Tc/(279.82+Tc)],   Tc < 0",11)
    d.add_heading("4.2 凝结 蒸发 冻结与融化",level=2)
    add_para(d,"一次时间步内先将产水加入蒸气库存，再执行饱和约束。若 mᵥ>mᵥ,sat，差额凝结为液水；若 mᵥ<mᵥ,sat，允许已有液水蒸发补足，蒸发量不超过液水库存。低于冰点时，液水按一阶动力学冻结，并受催化层孔隙最大储冰量限制。")
    add_eq(d,"Δmfreeze = min{mₗ[1−exp(−kfΔt)], ρice ε₀LCCL−mᵢ}",12)
    add_eq(d,"Δmmelt = min{mᵢ[1−exp(−kmΔt)], mᵢ},   T ≥ 273.15 K",13)
    add_para(d,"取 kf=0.35 s⁻¹、km=2.0 s⁻¹、ρice=920 kg·m⁻³。冻结释放的潜热为正，融化吸热为负：")
    add_eq(d,"Qlat = Lf Δmfreeze/Δt  或  −Lf Δmmelt/Δt",14)

    d.add_heading("5 冰堵塞对动力学和传质的耦合",level=1)
    add_para(d,"先把冰质量换算成阴极催化层内的绝对冰体积分数 εice，再除以初始孔隙率得到冰饱和度 sice。有效气相孔隙率保留一个数值下限，防止极限电流奇异。")
    add_eq(d,"εice = mᵢ/(ρiceLCCL),   sice = clip(εice/ε₀,0,0.999)",15)
    add_eq(d,"εeff = max(ε₀−εice, 0.015)",16)
    add_para(d,"Jiao 和 Li 的三维多相模型把催化反应源乘以未被液水和冰占据的比例，并用 1.5 次孔隙修正描述气体扩散。由于本题集中模型未单独追踪反应表面的液水覆盖，新版仅以冰饱和度修正交换电流。")
    add_eq(d,"a/a₀ = max(0.04, 1−sice)",17)
    add_eq(d,"DO₂,eff = DO₂,ref (T/Tref)¹·⁷⁵ εeff¹·⁵ MD",18)
    add_para(d,"上一版使用 (1−sice)³·⁵ 修正反应面积，同时又在扩散系数中计入冰占孔，造成冰影响的重复放大。残差诊断显示留出段电压误差与预测冰量的相关系数达到约 −0.93 至 −0.98。改为式（17）的线性覆盖后，两种温度的留出电压误差同时下降，因此保留该修正。")

    d.add_heading("6 膜水化与阳极暂态失水",level=1)
    d.add_heading("6.1 累计电荷驱动的膜水化",level=2)
    add_para(d,"附件没有入口湿度和膜内含水量。为保留产水对膜电导的主要作用，定义累计电荷 q，并用饱和增长函数近似有效膜含水量。这里 j 以 A·cm⁻²计，使 q 的单位为 C·cm⁻²。")
    add_eq(d,"qⁿ⁺¹ = qⁿ + 0.5(jⁿ+jⁿ⁺¹)Δt",19)
    add_eq(d,"λ(q)=λ₀+Δλ[1−exp(−q/qλ)]",20)
    add_eq(d,"κ(λ,T)=max{0.02, (0.5139λ−0.326)exp[1268(1/303.15−1/T)]}",21)
    add_para(d,"膜欧姆电阻为 LPEM/κ。式（20）是可辨识的低阶代理，并非空间分布膜水守恒方程；Δλ 和 qλ 因而只能解释为该工况下的有效水化参数。")
    d.add_heading("6.2 升载触发的失水记忆",level=2)
    add_para(d,"Yao 等指出，高电流冷启动可能首先出现阳极脱水，而较高初始膜含水量更容易转为阴极孔隙堵塞。为表示升载瞬间的阳极水亏缺，引入记忆状态 z。电流只在上升时激励 z；已有记忆按时间常数衰减；累计产水越多，后续升载的脱水激励越弱。")
    add_eq(d,"zⁿ⁺¹ = zⁿexp(−Δt/τd) + max(0,jⁿ⁺¹−jⁿ)exp(−qⁿ⁺¹/qd)",22)
    add_eq(d,"ηdry = Kd z",23)
    add_para(d,"取 τd=8 s、qd=0.6 C·cm⁻²为先验尺度，Kd 由训练数据拟合。曾尝试在中间时间窗重新选择 τd 和 qd，虽降低中间窗误差，但最终 30% 留出误差明显上升，因此未采用该组参数。")

    d.add_heading("7 电化学电压模型",level=1)
    d.add_heading("7.1 可逆电势与交换电流",level=2)
    add_eq(d,"Erev = 1.229−8.5×10⁻⁴(T−Tref)+RT/(2F) ln[pH₂(pO₂)¹ᐟ²]",24)
    add_eq(d,"j₀ = j₀,ref exp{−Ea/R(1/T−1/Tref)} · max(0.04,1−sice)",25)
    d.add_heading("7.2 四类电压损失",level=2)
    add_eq(d,"ηact = RT/(0.5F) asinh[j/(2j₀)]",26)
    add_eq(d,"ηohm = j(LPEM/κ + Rc),   Rc=0.01 Ω·cm²",27)
    add_eq(d,"jlim = 4F DO₂,eff cO₂/LGDL,   ηcon = −RT/(4F) ln(1−j/jlim)",28)
    add_eq(d,"V = Erev + ΔE − ηact − ηohm − ηcon − ηdry",29)
    add_para(d,"ΔE 吸收参考压力、开路电压与题面动力学常数之间不能由现有数据拆分的偏差。程序限制 j/jlim<0.999999，避免浓差项在极限电流附近数值发散。")

    d.add_heading("8 热源与多方程闭合",level=1)
    d.add_heading("8.1 反应热与潜热",level=2)
    add_para(d,"电化学放热采用热中性电压与端电压之差计算。式（29）的端电压降低会增加发热，从而升高温度；温度升高又提高膜电导和交换电流，构成电热正反馈。冰的形成同时释放凝固潜热，但又通过式（17）和式（18）提高电压损失，构成质量传递与电化学之间的负反馈。")
    add_eq(d,"Qrxn = j(Eth−V),   Eth=1.48 V",30)
    add_eq(d,"QCCL = Qrxn + Qlat",31)
    d.add_heading("8.2 单个时间步内的固定点迭代",level=2)
    add_para(d,"从 tⁿ 到 tⁿ⁺¹ 的求解采用分区固定点方法。水相库存每次迭代都从步初状态重新计算，防止同一时间步重复累计产水或冻结。具体顺序如下：")
    steps=[
        "读取 jⁿ、jⁿ⁺¹和 Δt，计算步内平均电流，更新累计电荷 q 和失水状态 z。",
        "以步初平均温度作为第 0 次猜测 T̄⁽⁰⁾；保存步初的 mᵥ、mₗ、mᵢ。",
        "用 T̄⁽ᵏ⁾计算饱和蒸气容量、凝结或蒸发、冻结或融化，得到候选水相状态和 Qlat。",
        "由候选冰质量计算 sice、有效反应面积和 DO₂,eff，再由 λ、z 和 T̄⁽ᵏ⁾计算步内电压 V⁽ᵏ⁾。",
        "以 V⁽ᵏ⁾形成 Qrxn，与 Qlat 相加后解七阶隐式热方程，得到 T₁⁽ᵏ⁺¹⁾…T₇⁽ᵏ⁺¹⁾。",
        "若 |T̄⁽ᵏ⁺¹⁾−T̄⁽ᵏ⁾|<10⁻⁸ K 且 |V⁽ᵏ⁾−V⁽ᵏ⁻¹⁾|<10⁻⁹ V，则接受该步；否则令 T̄⁽ᵏ⁾←T̄⁽ᵏ⁺¹⁾继续，最多 12 次。",
        "接受温度和水相状态后，以采样时刻 jⁿ⁺¹及更新温度重新计算输出电压，使输出与实验采样时点一致。"
    ]
    for i,s in enumerate(steps,1):
        p=d.add_paragraph(); p.paragraph_format.left_indent=Pt(20); p.paragraph_format.first_line_indent=Pt(-15); p.add_run(f"{i}. ").bold=True; p.add_run(s)
    add_para(d,"本数据上的每个时间步最多需要 4 次迭代，低于设置的 12 次上限。该闭合方式使温度、水相、电压和热源双向一致，同时保持每个子方程简单可检查。")

    d.add_heading("9 参数辨识与可辨识性",level=1)
    add_para(d,"每个温度工况共有 184 个观测点。前 128 点用于拟合，后 56 点留出。首先用实测电压驱动热模型，仅辨识统一热容倍率 sC；随后在实测温度下拟合电压参数；最后使用完整耦合模型进行训练段和留出段前向计算。参数优化采用带边界的阻尼 Gauss-Newton 方法。")
    add_table(d,["参数","数值","作用与性质"],[
        ("sC","1.02243","七层热容统一倍率，热数据拟合"),("j₀,ref","0.00765 A·m⁻²","有效交换电流"),
        ("Ea","66.27 kJ·mol⁻¹","表观活化能"),("MD","108.36","有效氧传质倍率"),("ΔE","0.1626 V","实验条件等效电势偏置"),
        ("Δλ","15.84","有效膜水化增量"),("qλ","1.689 C·cm⁻²","膜水化特征电荷"),("Kd","2.091 V/(A·cm⁻²)","失水记忆电压系数"),
        ("r","0.12","水滞留先验"),("kf","0.35 s⁻¹","冻结速率先验"),("τd  qd","8 s  0.6 C·cm⁻²","失水记忆先验")
    ],[1.05,1.65,3.3],[WD_ALIGN_PARAGRAPH.CENTER,WD_ALIGN_PARAGRAPH.CENTER,WD_ALIGN_PARAGRAPH.LEFT])
    add_para(d,"j₀,ref、MD、ΔE 和膜水化参数存在补偿关系，不能作为独立材料常数使用。尤其是 MD 数值较大，说明集中传质模型把未显式描述的流道、边界浓度和面积换算一并吸收到有效参数中。模型适合重现实验条件和比较相同装置下的策略，不应未经新数据校准直接外推到其他电池。")

    d.add_heading("10 与实验数据的对比",level=1)
    add_table(d,["工况与区段","电压 RMSE","电压平均相对误差","温度 RMSE","预测最大 εice"],[
        ("−20°C 训练","0.01133 V","1.56%","0.0270°C","0.02797"),("−20°C 留出","0.01990 V","2.47%","0.0586°C","0.06040"),
        ("−25°C 训练","0.00784 V","1.04%","0.0275°C","0.02884"),("−25°C 留出","0.01202 V","1.65%","0.1126°C","0.06211")
    ],[1.25,1.0,1.45,1.05,1.15],[WD_ALIGN_PARAGRAPH.CENTER]*5)
    add_table(d,["留出工况","上一版电压 RMSE","新版电压 RMSE","降幅"],[
        ("−20°C","0.02695 V","0.01990 V","26.2%"),("−25°C","0.02092 V","0.01202 V","42.6%")
    ],[1.25,1.55,1.55,1.0],[WD_ALIGN_PARAGRAPH.CENTER]*4)
    add_para(d,"数值检查覆盖所有时间步：水蒸气、液水和冰库存非负；冰体积分数不超过阴极催化层孔隙率；电压位于 0 至 1.5 V；输出不存在 NaN 或无穷值；步内耦合迭代不超过上限。留出结果表明线性冰覆盖修正对两个温度均有效，没有出现只改善单一工况的情况。")

    d.add_heading("11 模型适用范围与实验需求",level=1)
    add_para(d,"当前模型只由电流密度、端电压和平均温度约束。它能够用于本题问题一的电压和平均温升预测，也可用于比较相近工况下的升载方案。下列量仍属于模型内部预测：各层温度、液水与冰质量、冰饱和度、膜含水量和阳极失水状态。若要把这些状态参数解释为真实物理量，需要增加高频阻抗或膜电阻、出口水量、局部温度以及原位冰成像数据。")
    add_para(d,"冻结过程还可能包含过冷水的随机成核。现有确定性一阶冻结模型适合重复实验的平均趋势，但不描述小面积电池中单次冷启动的随机失效时间。后续若获得多次重复试验，可把 kf 改为温度相关的随机危险率，并用失效时间分布而非单条电压曲线校准。")

    d.add_heading("参考文献",level=1)
    refs=[
        "[1] Jiao K, Li X. Three-dimensional multiphase modeling of cold start processes in polymer electrolyte membrane fuel cells. Electrochimica Acta, 2009, 54: 6876–6891.",
        "[2] Yao L, Ma F, Peng J, Zhang J, Zhang Y, Shi J. Analysis of the Failure Modes in the Polymer Electrolyte Fuel Cell Cold-Start Process—Anode Dehydration or Cathode Pore Blockage. Energies, 2020, 13: 256.",
        "[3] Luo Y, Jiao K. Cold start of proton exchange membrane fuel cell. Progress in Energy and Combustion Science, 2017. DOI: 10.1016/j.pecs.2017.10.003.",
        "[4] Zhou Y, Luo Y, Yu S, Jiao K. Modeling of cold start processes and performance optimization for proton exchange membrane fuel cell stacks. Journal of Power Sources, 2014, 247: 738–748.",
        "[5] 陈政等. 一维 PEMFC 低温冷启动瞬态模型研究. 用户提供文献."
    ]
    for ref in refs:
        p=d.add_paragraph(ref); p.paragraph_format.left_indent=Pt(15); p.paragraph_format.first_line_indent=Pt(-15); p.paragraph_format.space_after=Pt(4)

    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    rr=footer.add_run("问题一  PEMFC 低温冷启动详细耦合模型"); rr.font.size=Pt(8); rr.font.color.rgb=RGBColor(90,90,90)
    d.core_properties.title="问题一 PEMFC 低温冷启动详细耦合模型"
    d.core_properties.subject="模型建立 数值求解 参数辨识 留出验证"
    d.core_properties.author=""
    d.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    build()
